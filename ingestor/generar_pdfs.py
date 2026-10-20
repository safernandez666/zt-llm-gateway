"""
Genera el corpus de prueba en formatos reales a partir de los markdown.

    docker compose run --rm ingestor python generar_pdfs.py

La mezcla no es casual: política de gastos y de viajes salen en PDF, que es
como llegan de verdad los documentos corporativos; vacaciones en DOCX; el
onboarding queda en markdown. Y el documento envenenado es un PDF — una
inyección escondida en un PDF que alguien subió al repositorio de políticas
es bastante más verosímil que una en un archivo de texto.
"""
import pathlib
import re

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

ORIGEN = pathlib.Path("/origen")
DESTINO = pathlib.Path("/documentos")

FORMATOS = {
    "politica-gastos.md": "pdf",
    "politica-viajes.md": "pdf",        # el envenenado
    "politica-vacaciones.md": "docx",
    "onboarding.md": "md",
}


def bloques(texto: str):
    """(nivel, contenido): nivel 0 es párrafo, 1..3 encabezado."""
    for linea in texto.splitlines():
        m = re.match(r"^(#{1,3})\s+(.*)", linea)
        if m:
            yield len(m.group(1)), m.group(2).strip()
        elif linea.strip():
            yield 0, linea.strip()


def a_pdf(texto: str, salida: pathlib.Path) -> None:
    estilos = getSampleStyleSheet()
    h = {1: estilos["Title"], 2: estilos["Heading2"], 3: estilos["Heading3"]}
    cuerpo = ParagraphStyle("cuerpo", parent=estilos["BodyText"], leading=15)
    doc = SimpleDocTemplate(str(salida), pagesize=A4,
                            leftMargin=2.5 * cm, rightMargin=2.5 * cm,
                            topMargin=2.5 * cm, bottomMargin=2.5 * cm)
    flujo = []
    for nivel, contenido in bloques(texto):
        # Los comentarios HTML del markdown no se arrastran al PDF: la
        # inyección tiene que verse como texto normal del documento.
        contenido = re.sub(r"<!--.*?-->", "", contenido).strip()
        if not contenido:
            continue
        flujo.append(Paragraph(contenido, h.get(nivel, cuerpo)))
        flujo.append(Spacer(1, 6 if nivel else 4))
    doc.build(flujo)


def a_docx(texto: str, salida: pathlib.Path) -> None:
    doc = Document()
    for nivel, contenido in bloques(texto):
        contenido = re.sub(r"<!--.*?-->", "", contenido).strip()
        if not contenido:
            continue
        if nivel:
            doc.add_heading(contenido, level=min(nivel, 4))
        else:
            doc.add_paragraph(contenido)
    doc.save(str(salida))


def main() -> None:
    DESTINO.mkdir(parents=True, exist_ok=True)
    for nombre, formato in FORMATOS.items():
        origen = ORIGEN / nombre
        if not origen.exists():
            print(f"  falta {origen}, lo salteo")
            continue
        texto = origen.read_text(encoding="utf-8")
        destino = DESTINO / (origen.stem + "." + formato)
        if formato == "pdf":
            a_pdf(texto, destino)
        elif formato == "docx":
            a_docx(texto, destino)
        else:
            destino.write_text(texto, encoding="utf-8")
        print(f"  {destino.name}  ({destino.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
